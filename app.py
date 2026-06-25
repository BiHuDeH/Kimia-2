import pandas as pd
import streamlit as st
import time
from io import BytesIO
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# --- تنظیمات صفحه ---
st.set_page_config(
    page_title="داشبورد مالی کیمیا",
    page_icon="💎",
    layout="wide",
    initial_sidebar_state="expanded"
)

# --- استایل‌دهی پیشرفته ---
def apply_modern_design():
    st.markdown("""
        <style>
            @import url('https://fonts.googleapis.com/css2?family=Vazirmatn:wght@100;300;400;500;700;900&display=swap');
            
            html, body, [class*="css"] {
                font-family: 'Vazirmatn', sans-serif !important;
                direction: rtl;
                text-align: right;
            }
            
            .stApp {
                background-color: #262b30 !important; 
                color: #f8f9fa !important;
            }
            
            h1, h2, h3, p, span, div {
                color: #e2e8f0 !important;
            }
            h1 {
                color: #60a5fa !important; 
                text-shadow: 1px 1px 3px rgba(0,0,0,0.4);
            }

            .stTabs [data-baseweb="tab-list"] {
                gap: 10px;
                background-color: #1e2429 !important;
                padding: 10px 10px 0 10px;
                border-radius: 12px 12px 0 0;
            }
            .stTabs [data-baseweb="tab"] {
                height: 50px;
                background-color: #343a40 !important;
                border-radius: 8px 8px 0 0;
                padding: 10px 20px;
                font-weight: bold;
                color: #adb5bd !important;
                border: 1px solid #495057 !important;
                border-bottom: none !important;
            }
            .stTabs [aria-selected="true"] {
                background-color: #3b82f6 !important;
                color: #ffffff !important;
            }

            [data-testid="stFileUploader"] {
                background-color: #343a40 !important;
                border: 2px dashed #6c757d !important;
                border-radius: 15px !important;
                padding: 20px !important;
                transition: all 0.3s ease;
            }
            [data-testid="stFileUploader"]:hover {
                border-color: #60a5fa !important;
                background-color: #3b4249 !important;
            }
            
            [data-testid="stFileUploader"] section {
                background-color: #495057 !important;
                border-radius: 8px;
            }
            
            [data-testid="stFileUploader"] small, [data-testid="stFileUploader"] div {
                color: #f8f9fa !important;
                font-size: 14px !important;
            }
            
            [data-testid="stFileUploader"] button[kind="secondary"] {
                background-color: #2563eb !important;
                border: none !important;
                position: relative;
                width: 160px;
                height: 40px;
                border-radius: 8px;
            }
            [data-testid="stFileUploader"] button[kind="secondary"] div, 
            [data-testid="stFileUploader"] button[kind="secondary"] span {
                display: none !important;
            }
            [data-testid="stFileUploader"] button[kind="secondary"]::after {
                content: "انتخاب فایل";
                color: white !important;
                position: absolute;
                top: 50%;
                left: 50%;
                transform: translate(-50%, -50%);
                font-family: 'Vazirmatn', sans-serif;
                font-size: 15px;
                font-weight: bold;
                display: block !important;
            }

            .stButton > button {
                width: 100%;
                border-radius: 12px;
                height: 55px;
                font-size: 17px;
                font-weight: 700;
                color: white !important;
                background: linear-gradient(135deg, #2563eb, #3b82f6) !important;
                border: none !important;
                box-shadow: 0 4px 15px rgba(37, 99, 235, 0.3);
                transition: transform 0.2s ease;
            }
            .stButton > button:hover {
                transform: translateY(-2px);
                box-shadow: 0 6px 20px rgba(37, 99, 235, 0.5);
            }

            .stDownloadButton > button {
                background: linear-gradient(135deg, #059669, #10b981) !important;
                box-shadow: 0 4px 15px rgba(16, 185, 129, 0.3) !important;
                height: 60px;
                font-size: 18px;
                border: 2px solid #047857 !important;
            }

            [data-testid="stDataFrame"] {
                background-color: #343a40 !important;
                border-radius: 12px;
                border: 1px solid #495057;
            }
            th {
                background-color: #212529 !important;
                color: #ffffff !important;
            }
            td {
                color: #e2e8f0 !important;
            }
            
        </style>
    """, unsafe_allow_html=True)

apply_modern_design()

# --- توابع کمکی ---

def to_persian_num(text):
    if not isinstance(text, str): text = str(text)
    translation = text.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")
    return text.translate(translation)

def clean_currency(val):
    if pd.isna(val) or val == "": return 0.0
    val_str = str(val).replace(",", "").replace("ریال", "").strip()
    try: return float(val_str)
    except ValueError: return 0.0

def find_header_row(df, keywords):
    for i in range(min(20, len(df))):
        row_values = df.iloc[i].astype(str).tolist()
        if any(str(keyword) in str(val) for val in row_values for keyword in keywords):
            return i
    return 0

def standardize_columns(df, col_map):
    df.columns = df.columns.astype(str)
    rename_dict = {}
    used_columns = set()
    for std_name, aliases in col_map.items():
        for col in df.columns:
            if col in used_columns: continue
            if any(str(alias) in str(col) for alias in aliases):
                rename_dict[col] = std_name
                used_columns.add(col)
                break 
    df = df.rename(columns=rename_dict)
    df = df.loc[:, ~df.columns.duplicated()]
    return df

def contains_all_words(text, phrase):
    if pd.isna(text): return False
    text = str(text)
    words = phrase.split()
    return all(word in text for word in words)

def match_any_phrase(text, phrases_list):
    for phrase in phrases_list:
        if contains_all_words(text, phrase):
            return True
    return False

def generate_styled_excel(df, sheet_name="Report"):
    output = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.sheet_view.rightToLeft = True 

    # تنظیمات فونت Calibri سایز ۱۱
    header_font = Font(bold=True, size=11, name='Calibri')
    regular_font = Font(size=11, name='Calibri')
    
    # تنظیمات وسط‌چین (عمودی و افقی)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    thin_border_side = Side(style="thin", color="333333")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    headers = list(df.columns)
    ws.append(headers)
    
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    
    for row in df.itertuples(index=False, name=None):
        ws.append(row)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.alignment = center_align
            cell.font = regular_font
            cell.border = thin_border
            if cell.column_letter != 'A':
                cell.number_format = '#,##0'

    # تنظیم هوشمند و متناسب عرض ستون‌ها
    for col in ws.columns:
        max_length = 0
        column_letter = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 6)
        if adjusted_width < 18:
            adjusted_width = 18
        ws.column_dimensions[column_letter].width = adjusted_width

    if ws.max_row >= 2:
        last_col = get_column_letter(ws.max_column)
        tab = Table(displayName=f"Table_{sheet_name.replace(' ', '_')}", ref=f"A1:{last_col}{ws.max_row}")
        style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        tab.tableStyleInfo = style
        ws.add_table(tab)

    wb.save(output)
    return output.getvalue()

# --- پردازشگر گزارش پاسارگاد ---
def process_pasargad(file):
    try:
        df = pd.read_excel(file, skiprows=2)
        expected_columns = [
            'Index', 'Branch Code', 'Branch', 'Date', 'Time', 
            'Document Number', 'Receipt Number', 'Check Number', 
            'Description', 'Withdrawal', 'Deposit', 'Balance', 'Notes'
        ]
        if len(df.columns) >= len(expected_columns):
            df.columns = expected_columns + list(df.columns[len(expected_columns):])
        else:
            return None, "ساختار فایل با الگوی استاندارد پاسارگاد مطابقت ندارد."

        if 'Date' not in df.columns: return None, "ستون Date یافت نشد."
        
        df = df.dropna(subset=['Date'])
        unique_dates = df['Date'].sort_values().unique()

        for col in ['Deposit', 'Withdrawal', 'Balance']:
            if col in df.columns:
                df[col] = pd.to_numeric(df[col].astype(str).str.replace(',', ''), errors='coerce').fillna(0)

        df['Description'] = df['Description'].astype(str)

        card_to_card_mask = df['Description'].str.contains("انتقال از", na=False)
        fee_mask = df['Description'].str.contains("کارمزد", na=False)
        
        # شرط ترکیبی (OR) برای برداشت روز
        cond1 = df['Description'].str.contains("انتقال وجه به سپرده 379.8000.10822179.1 به نام سپرده كوتاه مدت - مهران کلانتريان_سامانه بانکداری نوین", na=False)
        cond2 = df['Description'].str.contains("5022291034088359", na=False)
        daily_withdrawal_mask = cond1 | cond2
        
        snap_deposit_mask = df['Description'].apply(lambda x: contains_all_words(x, "غذا اطلس"))
        
        # ستون جدید: برداشت حضوری از بانک
        in_person_mask = df['Description'].str.contains("سند عملیات بانکی", na=False)

        card_to_card_sum = df[card_to_card_mask].groupby('Date')['Deposit'].sum().reindex(unique_dates, fill_value=0)
        fee_sum = df[fee_mask].groupby('Date')['Withdrawal'].sum().reindex(unique_dates, fill_value=0)
        daily_withdrawal_sum = df[daily_withdrawal_mask].groupby('Date')['Withdrawal'].sum().reindex(unique_dates, fill_value=0)
        snap_deposit_sum = df[snap_deposit_mask].groupby('Date')['Deposit'].sum().reindex(unique_dates, fill_value=0)
        in_person_sum = df[in_person_mask].groupby('Date')['Withdrawal'].sum().reindex(unique_dates, fill_value=0)
        end_of_day_balance = df.sort_values(['Date', 'Time']).groupby('Date')['Balance'].last().reindex(unique_dates, fill_value=0)

        report = pd.DataFrame(index=unique_dates)
        report.index.name = 'Date'
        report['Card_to_Card'] = card_to_card_sum
        report['Fee'] = fee_sum
        report['Daily_Withdrawal'] = daily_withdrawal_sum
        report['In_Person_Withdrawal'] = in_person_sum
        report['Snap_Deposit'] = snap_deposit_sum
        report['End_of_Day_Balance'] = end_of_day_balance
        report['Sales'] = report['Card_to_Card'] / 1.1
        report['Tax'] = report['Card_to_Card'] - report['Sales']

        report = report.reset_index()
        report.columns = ['تاریخ', 'کارت به کارت', 'کارمزد', 'برداشت روز', 'برداشت حضوری از بانک', 'واریزی اسنپ', 'مانده آخر روز', 'فروش', 'مالیات']
        
        final_order = ['تاریخ', 'کارت به کارت', 'فروش', 'مالیات', 'کارمزد', 'برداشت روز', 'برداشت حضوری از بانک', 'مانده آخر روز', 'واریزی اسنپ']
        
        return report[final_order], None

    except Exception as e:
        return None, str(e)

# --- پردازشگر گزارش کارآفرین ---
def process_karafrin(file):
    try:
        df_temp = pd.read_excel(file, header=None)
        header_row = find_header_row(df_temp, ["Date", "تاریخ"])
        df = pd.read_excel(file, header=header_row)
        
        col_map = {
            'Withdrawal': ['برداشت', 'بدهکار', 'Withdrawal', 'مبلغ برداشت'],
            'Balance': ['مانده', 'Balance', 'مانده حساب'],
            'Description': ['شرح', 'توضیحات', 'Description', 'شرح تراکنش'],
            'Date': ['تاریخ', 'Date', 'تاریخ تراکنش']
        }
        df = standardize_columns(df, col_map)
        
        req_cols = ['Date', 'Description', 'Withdrawal', 'Balance']
        missing = [c for c in req_cols if c not in df.columns]
        if missing: return None, f"ستون‌های الزامی یافت نشدند: {', '.join(missing)}"

        df = df.dropna(subset=['Date'])
        for col in ['Withdrawal', 'Balance']: df[col] = df[col].apply(clean_currency)
        df['Description'] = df['Description'].fillna("").astype(str)
        
        w_keywords = ["انتقال از", "برداشت از"]
        f_keywords = ["برداشت برای کارمزد", "دریافت کارمزد"]

        def calc_daily(group):
            w_sum = group[group['Description'].apply(lambda x: match_any_phrase(x, w_keywords))]['Withdrawal'].sum()
            f_sum = group[group['Description'].apply(lambda x: match_any_phrase(x, f_keywords))]['Withdrawal'].sum()
            first_bal = group['Balance'].iloc[0]
            return pd.Series({'برداشت روز': w_sum, 'کارمزد': f_sum, 'مانده روز': first_bal})

        result_df = df.groupby('Date', sort=False).apply(calc_daily).reset_index()
        result_df = result_df.rename(columns={'Date': 'تاریخ'})
        return result_df, None

    except Exception as e:
        return None, str(e)

# --- رابط کاربری اصلی ---
def main():
    st.markdown("""
        <div style="text-align: center; padding: 15px 0 30px 0;">
            <h1 style="font-size: 3.2em; margin-bottom: 5px;">داشبورد مالی کیمیا</h1>
            <p style="color: #adb5bd; font-size: 1.1em;">سیستم پردازش هوشمند صورتحساب‌ها</p>
        </div>
    """, unsafe_allow_html=True)

    tab1, tab2 = st.tabs(["🏦 گزارش مالی بانک پاسارگاد", "🏢 گزارش مالی بانک کارآفرین"])

    # ---------- تب پاسارگاد ----------
    with tab1:
        st.markdown("<br>", unsafe_allow_html=True)
        upl_pasargad = st.file_uploader("فایل اکسل پاسارگاد را اینجا بکشید و رها کنید (Drag & Drop) یا برای انتخاب کلیک کنید", type=["xlsx"], key="upl_pasargad")
        
        if upl_pasargad:
            st.success("✅ فایل با موفقیت دریافت شد.")
            
            if st.button("شروع پردازش داده‌های مربوط به فایل بانک پاسارگاد", key="btn_pasargad"):
                
                progress_text = "در حال تحلیل و استخراج داده‌ها از فایل پاسارگاد..."
                my_bar = st.progress(0, text=progress_text)
                for percent_complete in range(1, 101, 20):
                    time.sleep(0.1)
                    my_bar.progress(percent_complete, text=f"{progress_text} ({percent_complete}%)")

                res_pasargad, err_pasargad = process_pasargad(upl_pasargad)
                
                if err_pasargad:
                    my_bar.empty()
                    st.error(f"خطا در پردازش: {err_pasargad}")
                else:
                    my_bar.progress(100, text="پردازش با موفقیت به اتمام رسید! 💯")
                    time.sleep(0.5)
                    my_bar.empty()
                    
                    st.info("✨ فایل گزارش آماده دانلود است. لطفاً از دکمه زیر استفاده کنید:")
                    
                    excel_pasargad = generate_styled_excel(res_pasargad, "Pasargad Report")
                    
                    st.download_button(
                        "📥 دانلود فایل گزارش مالی بانک پاسارگاد",
                        excel_pasargad,
                        f"Pasargad_Report_{datetime.now().strftime('%Y%m%d')}.xlsx",
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="dl_pasargad"
                    )
                    
                    st.markdown("<hr style='border-color: #495057;'>", unsafe_allow_html=True)
                    st.markdown("#### پیش‌نمایش داده‌های استخراج شده:")
                    disp_pasargad = res_pasargad.copy()
                    for col in disp_pasargad.columns:
                        if col != 'تاریخ': disp_pasargad[col] = disp_pasargad[col].apply(lambda x: to_persian_num(f"{x:,.0f}"))
                    st.dataframe(disp_pasargad, use_container_width=True)

    # ---------- تب کارآفرین ----------
    with tab2:
        st.markdown("<br>", unsafe_allow_html=True)
        upl_karafrin = st.file_uploader("فایل اکسل کارآفرین را اینجا بکشید و رها کنید (Drag & Drop) یا برای انتخاب کلیک کنید", type=["xlsx"], key="upl_karafrin")
        
        if upl_karafrin:
            st.success("✅ فایل با موفقیت دریافت شد.")
            
            if st.button("شروع پردازش داده‌های مربوط به فایل بانک کارآفرین", key="btn_karafrin"):
                
                progress_text = "در حال تحلیل و استخراج داده‌ها از فایل کارآفرین..."
                my_bar = st.progress(0, text=progress_text)
                for percent_complete in range(1, 101, 20):
                    time.sleep(0.1)
                    my_bar.progress(percent_complete, text=f"{progress_text} ({percent_complete}%)")

                res_karafrin, err_karafrin = process_karafrin(upl_karafrin)
                
                if err_karafrin:
                    my_bar.empty()
                    st.error(f"خطا در پردازش: {err_karafrin}")
                else:
                    my_bar.progress(100, text="پردازش با موفقیت به اتمام رسید! 💯")
                    time.sleep(0.5)
                    my_bar.empty()
                    
                    st.info("✨ فایل گزارش آماده دانلود است. لطفاً از دکمه زیر استفاده کنید:")
                    
                    excel_karafrin = generate_styled_excel(res_karafrin, "Karafarin Report")
                    
                    st.download_button(
                        "📥 دانلود فایل گزارش مالی بانک کارآفرین",
                        excel_karafrin,
                        f"Karafarin_Report_{datetime.now().strftime('%Y%m%d')}.xlsx",
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="dl_karafrin"
                    )
                    
                    st.markdown("<hr style='border-color: #495057;'>", unsafe_allow_html=True)
                    st.markdown("#### پیش‌نمایش داده‌های استخراج شده:")
                    disp_karafrin = res_karafrin.copy()
                    for col in disp_karafrin.columns:
                        if col != 'تاریخ': disp_karafrin[col] = disp_karafrin[col].apply(lambda x: to_persian_num(f"{x:,.0f}"))
                    st.dataframe(disp_karafrin, use_container_width=True)

if __name__ == "__main__":
    main()
